"""
Evaluate T5 and Mac Foundation Model on 50 paragraph rephrasing examples.
Compare LLM evaluation results only. Clear progress display. Output: paragraph_rephrasing_evaluation.xlsx
"""

import os
import sys
import pandas as pd
from dotenv import load_dotenv

load_dotenv()

DATASET_CSV = "paragraph_rephrasing_dataset.csv"
OUTPUT_EXCEL = "paragraph_rephrasing_evaluation.xlsx"
NUM_EXAMPLES = 50


def progress(msg, current, total=NUM_EXAMPLES):
    """Print progress with clear formatting."""
    pct = 100 * current / total if total else 0
    print(f"  [{current:3d}/{total}] ({pct:5.1f}%) {msg}", flush=True)


def main():
    print("=" * 70)
    print("Paragraph Rephrasing Evaluation (50 examples, T5 vs Mac, LLM comparison)")
    print("=" * 70)

    # Ensure dataset exists
    if not os.path.isfile(DATASET_CSV):
        print("\nGenerating paragraph rephrasing dataset...")
        from dataset_paragraph_rephrasing import main as gen_main
        gen_main()
    print(f"\nStep 0: Loading dataset from {DATASET_CSV}")
    df = pd.read_csv(DATASET_CSV)
    df["All_References"] = df["All_References"].apply(
        lambda x: x.split("|") if isinstance(x, str) and "|" in x else ([x] if pd.notna(x) and x else [])
    )
    n = len(df)
    print(f"  Loaded {n} paragraph examples.")
    print()

    # T5 model
    print("Step 1: Running T5 model on paragraph rephrasing...")
    print("-" * 70)
    try:
        from model_loader import GrammarCorrectionModel
        t5_model = GrammarCorrectionModel()
    except Exception as e:
        print(f"  Failed to load T5: {e}")
        sys.exit(1)
    t5_results = []
    for idx, row in df.iterrows():
        input_text = row["Input_Text"]
        try:
            pred = t5_model.correct(input_text, max_length=512, num_beams=4)
        except Exception as e:
            print(f"  T5 error at {idx}: {e}")
            pred = input_text
        t5_results.append({"Index": idx, "Input_Text": input_text, "Ground_Truth": row["Ground_Truth"], "Predicted": pred, "All_References": "|".join(row["All_References"])})
        progress("T5 rephrasing", len(t5_results), n)
    t5_df = pd.DataFrame(t5_results)
    t5_df["Category"] = "Rephrasing"
    print("  T5 complete.")
    print()

    # Mac Foundation Model
    print("Step 2: Running Mac Foundation Model on paragraph rephrasing...")
    print("-" * 70)
    try:
        from mac_foundation_model_loader import MacFoundationModel, mac_foundation_model_available
        if not mac_foundation_model_available():
            raise RuntimeError("Mac Foundation Model not available")
        mac_model = MacFoundationModel()
    except Exception as e:
        print(f"  Mac model not available: {e}")
        mac_df = pd.DataFrame({"Index": t5_df["Index"], "Mac_Predicted": [""] * len(t5_df)})
    else:
        mac_results = []
        for idx, row in df.iterrows():
            input_text = row["Input_Text"]
            try:
                pred = mac_model.correct(input_text, category="Rephrasing")
            except Exception as e:
                print(f"  Mac error at {idx}: {e}")
                pred = input_text
            mac_results.append({"Index": idx, "Mac_Predicted": pred})
            progress("Mac rephrasing", len(mac_results), n)
        mac_df = pd.DataFrame(mac_results)
        if hasattr(mac_model, "close"):
            mac_model.close()
        print("  Mac complete.")
    print()

    # Merge Mac predictions into one table
    eval_df = t5_df.merge(mac_df[["Index", "Mac_Predicted"]], on="Index", how="left")
    eval_df["Mac_Predicted"] = eval_df["Mac_Predicted"].fillna("")

    # LLM evaluation on T5
    print("Step 3: LLM evaluation of T5 rephrasing outputs...")
    print("-" * 70)
    from llm_evaluator import LLMEvaluator
    llm_eval = LLMEvaluator()
    eval_df["T5_LLM_Score"] = pd.NA
    eval_df["T5_LLM_Explanation"] = ""
    eval_df["T5_LLM_Explanation"] = eval_df["T5_LLM_Explanation"].astype(object)
    for i, idx in enumerate(eval_df["Index"]):
        row = eval_df.loc[eval_df["Index"] == idx].iloc[0]
        refs = row["All_References"].split("|") if isinstance(row["All_References"], str) and "|" in row["All_References"] else []
        r = llm_eval.evaluate_rephrasing(row["Input_Text"], row["Ground_Truth"], row["Predicted"], refs)
        eval_df.loc[eval_df["Index"] == idx, "T5_LLM_Score"] = r["score"]
        eval_df.loc[eval_df["Index"] == idx, "T5_LLM_Explanation"] = r["explanation"]
        progress("LLM evaluating T5", i + 1, n)
    print("  T5 LLM evaluation complete.")
    print()

    # LLM evaluation on Mac
    print("Step 4: LLM evaluation of Mac Foundation Model rephrasing outputs...")
    print("-" * 70)
    eval_df["Mac_LLM_Score"] = pd.NA
    eval_df["Mac_LLM_Explanation"] = ""
    eval_df["Mac_LLM_Explanation"] = eval_df["Mac_LLM_Explanation"].astype(object)
    for i, idx in enumerate(eval_df["Index"]):
        row = eval_df.loc[eval_df["Index"] == idx].iloc[0]
        refs = row["All_References"].split("|") if isinstance(row["All_References"], str) and "|" in row["All_References"] else []
        r = llm_eval.evaluate_rephrasing(row["Input_Text"], row["Ground_Truth"], row["Mac_Predicted"], refs)
        eval_df.loc[eval_df["Index"] == idx, "Mac_LLM_Score"] = r["score"]
        eval_df.loc[eval_df["Index"] == idx, "Mac_LLM_Explanation"] = r["explanation"]
        progress("LLM evaluating Mac", i + 1, n)
    print("  Mac LLM evaluation complete.")
    print()

    # Summary: LLM comparison only
    print("Step 5: Building LLM comparison summary...")
    t5_avg = eval_df["T5_LLM_Score"].mean()
    mac_avg = eval_df["Mac_LLM_Score"].mean()
    t5_4plus = (eval_df["T5_LLM_Score"] >= 4).sum()
    mac_4plus = (eval_df["Mac_LLM_Score"] >= 4).sum()
    t5_5 = (eval_df["T5_LLM_Score"] == 5).sum()
    mac_5 = (eval_df["Mac_LLM_Score"] == 5).sum()
    print(f"  T5  - Avg LLM score: {t5_avg:.2f}, Score >= 4: {t5_4plus}, Score 5: {t5_5}")
    print(f"  Mac - Avg LLM score: {mac_avg:.2f}, Score >= 4: {mac_4plus}, Score 5: {mac_5}")
    print()

    # Write Excel
    print("Step 6: Writing Excel report...")
    print("-" * 70)
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill

    input_sheet = df[["Index", "Category", "Input_Text", "Ground_Truth", "All_References"]].copy()
    input_sheet["All_References"] = input_sheet["All_References"].apply(lambda x: "|".join(x) if isinstance(x, list) else str(x))

    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        input_sheet.to_excel(writer, sheet_name="Input_Dataset", index=False)
        t5_df[["Index", "Input_Text", "Ground_Truth", "Predicted", "All_References"]].to_excel(writer, sheet_name="T5_Results", index=False)
        eval_df[["Index", "Input_Text", "Ground_Truth", "Mac_Predicted", "All_References"]].to_excel(writer, sheet_name="Mac_Results", index=False)
        llm_comp = eval_df[["Index", "Input_Text", "Ground_Truth", "Predicted", "Mac_Predicted", "T5_LLM_Score", "Mac_LLM_Score", "T5_LLM_Explanation", "Mac_LLM_Explanation"]].copy()
        llm_comp.columns = ["Index", "Input_Text", "Ground_Truth", "T5_Predicted", "Mac_Predicted", "T5_LLM_Score", "Mac_LLM_Score", "T5_LLM_Explanation", "Mac_LLM_Explanation"]
        llm_comp.to_excel(writer, sheet_name="LLM_Comparison", index=False)
        summary = pd.DataFrame([
            {"Metric": "Avg_LLM_Score", "T5": t5_avg, "Mac": mac_avg},
            {"Metric": "Score_4+_Count", "T5": t5_4plus, "Mac": mac_4plus},
            {"Metric": "Score_5_Count", "T5": t5_5, "Mac": mac_5},
        ])
        summary.to_excel(writer, sheet_name="LLM_Summary", index=False)
    try:
        wb = load_workbook(OUTPUT_EXCEL)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
        wb.save(OUTPUT_EXCEL)
    except Exception:
        pass
    print(f"  Saved to {OUTPUT_EXCEL}")
    print()
    print("=" * 70)
    print("Evaluation complete.")
    print("=" * 70)


if __name__ == "__main__":
    main()
