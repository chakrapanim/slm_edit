"""
Simple script to add All_References column to Excel file
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def format_excel_report(filename):
    """Format the Excel report with styling"""
    wb = load_workbook(filename)
    
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    category_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                if cell.column == 2:
                    cell.fill = category_fill
                    cell.font = Font(bold=True)
        
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        ws.freeze_panes = 'A2'
    
    wb.save(filename)

# Load dataset
dataset_df = pd.read_csv('test_dataset_expanded.csv')

# Load Excel sheets
detailed_df = pd.read_excel('model_evaluation_expanded_report.xlsx', sheet_name='Detailed_Results', engine='openpyxl')
llm_df = pd.read_excel('model_evaluation_expanded_report.xlsx', sheet_name='LLM_Evaluation', engine='openpyxl')

# Merge All_References
detailed_df = detailed_df.merge(
    dataset_df[['Category', 'Input_Text', 'All_References']],
    on=['Category', 'Input_Text'],
    how='left'
)

llm_df = llm_df.merge(
    dataset_df[['Category', 'Input_Text', 'All_References']],
    on=['Category', 'Input_Text'],
    how='left'
)

# Reorder columns
cols = list(detailed_df.columns)
if 'All_References' in cols:
    cols.remove('All_References')
    pred_idx = cols.index('Predicted') if 'Predicted' in cols else len(cols)
    cols.insert(pred_idx + 1, 'All_References')
    detailed_df = detailed_df[cols]

llm_cols = list(llm_df.columns)
if 'All_References' in llm_cols:
    llm_cols.remove('All_References')
    pred_idx = llm_cols.index('Predicted') if 'Predicted' in llm_cols else len(llm_cols)
    llm_cols.insert(pred_idx + 1, 'All_References')
    llm_df = llm_df[llm_cols]

# Read all other sheets
xl_file = pd.ExcelFile('model_evaluation_expanded_report.xlsx', engine='openpyxl')
all_sheets = {}
for sheet_name in xl_file.sheet_names:
    if sheet_name == 'Detailed_Results':
        all_sheets[sheet_name] = detailed_df
    elif sheet_name == 'LLM_Evaluation':
        all_sheets[sheet_name] = llm_df
    else:
        all_sheets[sheet_name] = pd.read_excel('model_evaluation_expanded_report.xlsx', sheet_name=sheet_name, engine='openpyxl')

# Write all sheets
with pd.ExcelWriter('model_evaluation_expanded_report.xlsx', engine='openpyxl') as writer:
    for sheet_name, df in all_sheets.items():
        df.to_excel(writer, sheet_name=sheet_name, index=False)

# Format
format_excel_report('model_evaluation_expanded_report.xlsx')

print("✓ Excel file updated with All_References column!")
