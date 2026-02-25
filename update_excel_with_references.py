"""
Script to update the existing Excel file to include All_References column
This adds the 4 reference variants for rephrasing examples
"""

import pandas as pd
import openpyxl
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def format_excel_report(filename):
    """Format the Excel report with styling"""
    wb = load_workbook(filename)
    
    # Define styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    category_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Format header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        # Format data rows
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                
                # Highlight category column
                if cell.column == 2:  # Category column
                    cell.fill = category_fill
                    cell.font = Font(bold=True)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        ws.freeze_panes = 'A2'
    
    wb.save(filename)
    print(f"Excel report formatted and saved to {filename}")

def main():
    print("=" * 80)
    print("Updating Excel file with All_References column")
    print("=" * 80)
    print()
    
    # Load the dataset to get All_References
    print("Loading dataset...")
    dataset_df = pd.read_csv('test_dataset_expanded.csv')
    dataset_df['All_References'] = dataset_df['All_References'].apply(
        lambda x: x if isinstance(x, str) else ''
    )
    print(f"Loaded {len(dataset_df)} examples from dataset")
    print()
    
    # Load existing Excel file
    print("Loading existing Excel file...")
    excel_file = 'model_evaluation_expanded_report.xlsx'
    
    # Read all sheets
    detailed_df = pd.read_excel(excel_file, sheet_name='Detailed_Results')
    llm_df = pd.read_excel(excel_file, sheet_name='LLM_Evaluation') if 'LLM_Evaluation' in pd.ExcelFile(excel_file).sheet_names else None
    
    print(f"Loaded {len(detailed_df)} rows from Detailed_Results")
    if llm_df is not None:
        print(f"Loaded {len(llm_df)} rows from LLM_Evaluation")
    print()
    
    # Merge All_References from dataset
    print("Adding All_References column...")
    # Match by Input_Text and Category
    detailed_df = detailed_df.merge(
        dataset_df[['Category', 'Input_Text', 'All_References']],
        on=['Category', 'Input_Text'],
        how='left'
    )
    
    # Reorder columns to put All_References after Predicted
    cols = list(detailed_df.columns)
    if 'All_References' in cols:
        cols.remove('All_References')
        # Find index of 'Predicted'
        pred_idx = cols.index('Predicted') if 'Predicted' in cols else len(cols)
        cols.insert(pred_idx + 1, 'All_References')
        detailed_df = detailed_df[cols]
    
    if llm_df is not None:
        llm_df = llm_df.merge(
            dataset_df[['Category', 'Input_Text', 'All_References']],
            on=['Category', 'Input_Text'],
            how='left'
        )
        # Reorder LLM columns too
        llm_cols = list(llm_df.columns)
        if 'All_References' in llm_cols:
            llm_cols.remove('All_References')
            pred_idx = llm_cols.index('Predicted') if 'Predicted' in llm_cols else len(llm_cols)
            llm_cols.insert(pred_idx + 1, 'All_References')
            llm_df = llm_df[llm_cols]
    
    print("✓ All_References column added")
    print()
    
    # Save updated Excel file
    print("Saving updated Excel file...")
    
    # First, read all existing sheets
    excel_sheets = {}
    xl_file = pd.ExcelFile(excel_file)
    for sheet_name in xl_file.sheet_names:
        if sheet_name == 'Detailed_Results':
            excel_sheets[sheet_name] = detailed_df
        elif sheet_name == 'LLM_Evaluation' and llm_df is not None:
            excel_sheets[sheet_name] = llm_df
        else:
            excel_sheets[sheet_name] = pd.read_excel(excel_file, sheet_name=sheet_name)
    
    # Now write all sheets
    with pd.ExcelWriter(excel_file, engine='openpyxl') as writer:
        for sheet_name, df in excel_sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
    
    # Format the file
    format_excel_report(excel_file)
    
    print()
    print("=" * 80)
    print("Update Complete!")
    print("=" * 80)
    print(f"\nUpdated file: {excel_file}")
    print("\nThe All_References column has been added showing all 4 reference variants")
    print("for rephrasing examples (and single reference for typos/grammar).")
    
    # Show sample
    print("\nSample rephrasing entry with references:")
    rephrasing = detailed_df[detailed_df['Category'] == 'Rephrasing']
    if len(rephrasing) > 0:
        sample = rephrasing.iloc[0]
        print(f"  Input: {sample['Input_Text']}")
        print(f"  Ground Truth: {sample['Ground_Truth']}")
        print(f"  All References: {sample['All_References']}")

if __name__ == "__main__":
    main()
