"""
Build final comparison Excel: T5 results vs Mac results.
- T5: from model_evaluation_expanded_report.xlsx (all categories).
- Mac: Typo + Grammar from mac_foundation_model_evaluation.xlsx (previous run);
       Rephrasing from mac_rephrasing_only_results.xlsx (new rephrasing prompt).
Output: final_comparison.xlsx
"""

import os
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

T5_REPORT = "model_evaluation_expanded_report.xlsx"
MAC_FULL_REPORT = "mac_foundation_model_evaluation.xlsx"
MAC_REPHRASING_REPORT = "mac_rephrasing_only_results.xlsx"
OUTPUT_EXCEL = "final_comparison.xlsx"


def _high_quality_count(df, category, dist_col="Normalized_Edit_Distance", threshold=0.8):
    """Count rows in df with Category == category and dist_col > threshold."""
    if df is None or df.empty or dist_col not in df.columns:
        return None
    cat_df = df[df["Category"] == category]
    if cat_df.empty:
        return None
    return len(cat_df[cat_df[dist_col] > threshold])


def _llm_avg(df, category):
    """Average LLM_Score for category in df."""
    if df is None or "LLM_Score" not in df.columns:
        return None
    cat_df = df[(df["Category"] == category) & (df["LLM_Score"].notna())]
    return cat_df["LLM_Score"].mean() if len(cat_df) > 0 else None


def main():
    print("=" * 70)
    print("Building final comparison: T5 vs Mac (Typos/Grammar previous, Rephrasing new prompt)")
    print("=" * 70)

    # Load T5
    print("\nLoading T5 results...")
    try:
        t5_cat = pd.read_excel(T5_REPORT, sheet_name="Category_Analysis", engine="openpyxl")
    except Exception as e:
        print(f"  Error loading T5: {e}")
        t5_cat = None

    # Load Mac Typo + Grammar (previous full run)
    print("Loading Mac Typo/Grammar (previous run)...")
    try:
        mac_summary = pd.read_excel(MAC_FULL_REPORT, sheet_name="Mac_Summary_Metrics", engine="openpyxl")
        mac_detailed = pd.read_excel(MAC_FULL_REPORT, sheet_name="Mac_Detailed_Results", engine="openpyxl")
        mac_llm_summary = pd.read_excel(MAC_FULL_REPORT, sheet_name="Mac_LLM_Summary", engine="openpyxl")
        mac_typo_grammar = mac_detailed[mac_detailed["Category"].isin(["Typo", "Grammar"])].copy()
        mac_summary_tg = mac_summary[mac_summary["Category"].isin(["Typo", "Grammar"])]
    except Exception as e:
        print(f"  Error loading Mac full report: {e}")
        mac_summary_tg = None
        mac_typo_grammar = None
        mac_llm_summary = None

    # Load Mac Rephrasing (new rephrasing-only run)
    print("Loading Mac Rephrasing (new prompt run)...")
    if not os.path.isfile(MAC_REPHRASING_REPORT):
        print(f"  {MAC_REPHRASING_REPORT} not found.")
        print("  Run first: python3 run_mac_rephrasing_only.py")
        return
    else:
        try:
            rephrase_summary = pd.read_excel(MAC_REPHRASING_REPORT, sheet_name="Rephrasing_Summary", engine="openpyxl")
            rephrase_detailed = pd.read_excel(MAC_REPHRASING_REPORT, sheet_name="Rephrasing_Detailed", engine="openpyxl")
        except Exception as e:
            print(f"  Error loading Mac rephrasing report: {e}")
            rephrase_summary = None
            rephrase_detailed = None

    def t5_val(cat, key):
        if t5_cat is None or t5_cat.empty:
            return None
        row = t5_cat[t5_cat["Category"] == cat]
        return row.iloc[0].get(key) if not row.empty else None

    def mac_typo_grammar_val(cat, key):
        if mac_summary_tg is None or mac_summary_tg.empty:
            return None
        row = mac_summary_tg[mac_summary_tg["Category"] == cat]
        if row.empty:
            return None
        col = "Avg_BLEU_Score" if key == "Avg_BLEU" else key
        return row.iloc[0].get(col)

    def mac_rephrase_val(key):
        if rephrase_summary is None or rephrase_summary.empty:
            return None
        r = rephrase_summary.iloc[0]
        col = "Avg_BLEU_Score" if key == "Avg_BLEU" else key
        return r.get(col)

    rows = []
    for cat in ["Typo", "Grammar", "Rephrasing"]:
        if cat == "Rephrasing":
            mac_em = mac_rephrase_val("Exact_Match_Rate")
            mac_bleu = mac_rephrase_val("Avg_BLEU_Score")
            mac_rougel = mac_rephrase_val("Avg_ROUGEL")
            mac_llm = _llm_avg(rephrase_detailed, "Rephrasing")
            mac_hq = _high_quality_count(rephrase_detailed, "Rephrasing")
        else:
            mac_em = mac_typo_grammar_val(cat, "Exact_Match_Rate")
            mac_bleu = mac_typo_grammar_val(cat, "Avg_BLEU_Score")
            mac_rougel = mac_typo_grammar_val(cat, "Avg_ROUGEL")
            mac_llm = _llm_avg(mac_typo_grammar, cat) if mac_typo_grammar is not None else None
            mac_hq = _high_quality_count(mac_typo_grammar, cat) if mac_typo_grammar is not None else None

        rows.append({
            "Category": cat,
            "T5_Exact_Match_Rate": t5_val(cat, "Exact_Match_Rate"),
            "Mac_Exact_Match_Rate": mac_em,
            "T5_Avg_BLEU": t5_val(cat, "Avg_BLEU"),
            "Mac_Avg_BLEU": mac_bleu,
            "T5_Avg_ROUGEL": t5_val(cat, "Avg_ROUGEL"),
            "Mac_Avg_ROUGEL": mac_rougel,
            "T5_Avg_LLM_Score": t5_val(cat, "Avg_LLM_Score"),
            "Mac_Avg_LLM_Score": mac_llm,
            "T5_High_Quality_Count": t5_val(cat, "High_Quality_Corrections"),
            "Mac_High_Quality_Count": mac_hq,
        })

    # Overall row (averages / aggregates)
    def overall_mac_exact():
        if mac_summary_tg is None or rephrase_summary is None:
            return None
        tg = mac_summary_tg
        em_t = tg[tg["Category"] == "Typo"]["Exact_Matches"].iloc[0] if len(tg[tg["Category"] == "Typo"]) else 0
        em_g = tg[tg["Category"] == "Grammar"]["Exact_Matches"].iloc[0] if len(tg[tg["Category"] == "Grammar"]) else 0
        em_r = int(rephrase_summary.iloc[0].get("Exact_Matches", 0))
        return (em_t + em_g + em_r) / 300.0

    def overall_mac_bleu():
        if mac_summary_tg is None or rephrase_summary is None:
            return None
        tg = mac_summary_tg
        b_t = tg[tg["Category"] == "Typo"]["Avg_BLEU_Score"].iloc[0] if len(tg[tg["Category"] == "Typo"]) else 0
        b_g = tg[tg["Category"] == "Grammar"]["Avg_BLEU_Score"].iloc[0] if len(tg[tg["Category"] == "Grammar"]) else 0
        b_r = rephrase_summary.iloc[0].get("Avg_BLEU_Score", 0)
        return (b_t + b_g + b_r) / 3.0

    def overall_mac_rougel():
        if mac_summary_tg is None or rephrase_summary is None:
            return None
        tg = mac_summary_tg
        r_t = tg[tg["Category"] == "Typo"]["Avg_ROUGEL"].iloc[0] if len(tg[tg["Category"] == "Typo"]) else 0
        r_g = tg[tg["Category"] == "Grammar"]["Avg_ROUGEL"].iloc[0] if len(tg[tg["Category"] == "Grammar"]) else 0
        r_r = rephrase_summary.iloc[0].get("Avg_ROUGEL", 0)
        return (r_t + r_g + r_r) / 3.0

    rows.append({
        "Category": "Overall",
        "T5_Exact_Match_Rate": t5_val("Overall", "Exact_Match_Rate"),
        "Mac_Exact_Match_Rate": overall_mac_exact(),
        "T5_Avg_BLEU": t5_val("Overall", "Avg_BLEU"),
        "Mac_Avg_BLEU": overall_mac_bleu(),
        "T5_Avg_ROUGEL": t5_val("Overall", "Avg_ROUGEL"),
        "Mac_Avg_ROUGEL": overall_mac_rougel(),
        "T5_Avg_LLM_Score": t5_val("Overall", "Avg_LLM_Score"),
        "Mac_Avg_LLM_Score": None,  # could aggregate from detailed if needed
        "T5_High_Quality_Count": t5_val("Overall", "High_Quality_Corrections"),
        "Mac_High_Quality_Count": (
            (_high_quality_count(mac_typo_grammar, "Typo") or 0) +
            (_high_quality_count(mac_typo_grammar, "Grammar") or 0) +
            (_high_quality_count(rephrase_detailed, "Rephrasing") or 0)
        ) or None,
    })

    comparison_df = pd.DataFrame(rows)

    # LLM Evaluation Comparison sheet: Avg LLM score + 4+ count + 5 count for T5 vs Mac
    def mac_llm_4_plus(cat):
        if cat == "Rephrasing":
            if rephrase_detailed is None or "LLM_Score" not in rephrase_detailed.columns:
                return None
            return int((rephrase_detailed["LLM_Score"] >= 4).sum())
        if mac_llm_summary is None or mac_llm_summary.empty:
            return None
        row = mac_llm_summary[mac_llm_summary["Category"] == cat]
        if row.empty:
            return None
        r = row.iloc[0]
        return int(r.get("Score_5_Count", 0) + r.get("Score_4_Count", 0))

    def mac_llm_5_count(cat):
        if cat == "Rephrasing":
            if rephrase_detailed is None or "LLM_Score" not in rephrase_detailed.columns:
                return None
            return int((rephrase_detailed["LLM_Score"] == 5).sum())
        if mac_llm_summary is None or mac_llm_summary.empty:
            return None
        row = mac_llm_summary[mac_llm_summary["Category"] == cat]
        if row.empty:
            return None
        return int(row.iloc[0].get("Score_5_Count", 0))

    llm_rows = []
    for cat in ["Typo", "Grammar", "Rephrasing"]:
        mac_avg = _llm_avg(rephrase_detailed, "Rephrasing") if cat == "Rephrasing" else _llm_avg(mac_typo_grammar, cat)
        llm_rows.append({
            "Category": cat,
            "T5_Avg_LLM_Score": t5_val(cat, "Avg_LLM_Score"),
            "Mac_Avg_LLM_Score": mac_avg,
            "T5_Score_4+_Count": t5_val(cat, "LLM_Score_4+"),
            "Mac_Score_4+_Count": mac_llm_4_plus(cat),
            "T5_Score_5_Count": t5_val(cat, "LLM_Score_5"),
            "Mac_Score_5_Count": mac_llm_5_count(cat),
        })
    # Overall row (sums for counts, average for avg score)
    t5_overall_avg = None
    mac_overall_avg = None
    if t5_cat is not None and not t5_cat.empty:
        overall = t5_cat[t5_cat["Category"] == "Overall"]
        if not overall.empty:
            t5_overall_avg = overall.iloc[0].get("Avg_LLM_Score")
    if llm_rows:
        mac_avgs = [r["Mac_Avg_LLM_Score"] for r in llm_rows if r["Mac_Avg_LLM_Score"] is not None]
        mac_overall_avg = sum(mac_avgs) / len(mac_avgs) if mac_avgs else None
    llm_rows.append({
        "Category": "Overall",
        "T5_Avg_LLM_Score": t5_overall_avg,
        "Mac_Avg_LLM_Score": mac_overall_avg,
        "T5_Score_4+_Count": (
            (llm_rows[0].get("T5_Score_4+_Count") or 0) +
            (llm_rows[1].get("T5_Score_4+_Count") or 0) +
            (llm_rows[2].get("T5_Score_4+_Count") or 0)
        ) if len(llm_rows) >= 3 else None,
        "Mac_Score_4+_Count": (
            (llm_rows[0].get("Mac_Score_4+_Count") or 0) +
            (llm_rows[1].get("Mac_Score_4+_Count") or 0) +
            (llm_rows[2].get("Mac_Score_4+_Count") or 0)
        ) if len(llm_rows) >= 3 else None,
        "T5_Score_5_Count": (
            (llm_rows[0].get("T5_Score_5_Count") or 0) +
            (llm_rows[1].get("T5_Score_5_Count") or 0) +
            (llm_rows[2].get("T5_Score_5_Count") or 0)
        ) if len(llm_rows) >= 3 else None,
        "Mac_Score_5_Count": (
            (llm_rows[0].get("Mac_Score_5_Count") or 0) +
            (llm_rows[1].get("Mac_Score_5_Count") or 0) +
            (llm_rows[2].get("Mac_Score_5_Count") or 0)
        ) if len(llm_rows) >= 3 else None,
    })
    llm_comparison_df = pd.DataFrame(llm_rows)

    # Sources sheet
    sources = pd.DataFrame([
        {"Sheet": "Comparison_Table", "Description": "T5 vs Mac metrics by category. Mac = Typo/Grammar from previous run, Rephrasing from new rephrasing prompt."},
        {"Sheet": "LLM_Evaluation_Comparison", "Description": "LLM scores (1-5): average and counts of 4+ and 5 by category for T5 vs Mac."},
        {"Sheet": "Sources", "Description": "T5: model_evaluation_expanded_report.xlsx. Mac Typo/Grammar: mac_foundation_model_evaluation.xlsx. Mac Rephrasing: mac_rephrasing_only_results.xlsx."},
    ])

    # Write Excel
    print("\nWriting final comparison Excel...")
    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        comparison_df.to_excel(writer, sheet_name="Comparison_Table", index=False)
        llm_comparison_df.to_excel(writer, sheet_name="LLM_Evaluation_Comparison", index=False)
        sources.to_excel(writer, sheet_name="Sources", index=False)

    # Format
    try:
        wb = load_workbook(OUTPUT_EXCEL)
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
        wb.save(OUTPUT_EXCEL)
    except Exception:
        pass

    print(f"  Saved to {OUTPUT_EXCEL}")
    print("\nComparison table:")
    print(comparison_df.to_string(index=False))
    print("\nLLM Evaluation Comparison:")
    print(llm_comparison_df.to_string(index=False))
    print("\nDone.")


if __name__ == "__main__":
    main()
