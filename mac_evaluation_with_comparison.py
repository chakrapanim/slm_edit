"""
Evaluate Mac Foundation Model (Apple Intelligence) on the expanded test dataset,
then compare with T5-base-grammar-correction and write a single Excel report.
Output: mac_foundation_model_evaluation.xlsx
"""

import pandas as pd
import os
import json
from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

# Optional Mac Foundation Model
try:
    from mac_foundation_model_loader import MacFoundationModel, mac_foundation_model_available
except Exception:
    MacFoundationModel = None
    mac_foundation_model_available = lambda: False

from evaluator import ModelEvaluator
from llm_evaluator import LLMEvaluator

T5_REPORT_PATH = "model_evaluation_expanded_report.xlsx"
TEST_DATASET_PATH = "test_dataset_expanded.csv"
OUTPUT_EXCEL = "mac_foundation_model_evaluation.xlsx"
PROGRESS_FILE = "mac_eval_progress.json"
TOTAL_EXAMPLES = 300


def _write_progress(phase, completed, total=TOTAL_EXAMPLES):
    """Write progress to a JSON file so you can check how many evaluations are done."""
    try:
        with open(PROGRESS_FILE, "w") as f:
            json.dump({
                "phase": phase,
                "completed": completed,
                "total": total,
                "message": f"{completed}/{total}",
            }, f, indent=2)
    except Exception:
        pass


def evaluate_with_multiple_references(predicted, all_references, evaluator, max_refs=4):
    if not all_references:
        return {}
    refs_to_eval = all_references[:max_refs] if len(all_references) > max_refs else all_references
    best_metrics = None
    best_score = -1
    for ref in refs_to_eval:
        metrics = evaluator.calculate_all_metrics(predicted, ref)
        score = metrics["normalized_edit_distance"]
        if score > best_score:
            best_score = score
            best_metrics = metrics
    all_metrics = []
    for ref in refs_to_eval:
        all_metrics.append(evaluator.calculate_all_metrics(predicted, ref))
    avg_metrics = {}
    if all_metrics:
        for key in all_metrics[0].keys():
            avg_metrics[f"avg_{key}"] = sum(m[key] for m in all_metrics) / len(all_metrics)
    result_metrics = (best_metrics or {}).copy()
    result_metrics.update(avg_metrics)
    result_metrics["best_ref_score"] = best_score
    return result_metrics


def format_excel_report(filename):
    wb = load_workbook(filename)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    category_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                if cell.column == 2:
                    cell.fill = category_fill
                    cell.font = Font(bold=True)
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)
        ws.freeze_panes = "A2"
    wb.save(filename)
    print(f"Excel report formatted and saved to {filename}")


def load_t5_results():
    """Load T5 metrics from existing expanded report."""
    if not os.path.exists(T5_REPORT_PATH):
        return None, None, None
    try:
        detailed = pd.read_excel(T5_REPORT_PATH, sheet_name="Detailed_Results", engine="openpyxl")
        detailed.columns = [c.strip() for c in detailed.columns]
        col_map = {
            "Exact_Match": "exact_match",
            "Char_Accuracy": "char_accuracy",
            "Word_Accuracy": "word_accuracy",
            "Normalized_Edit_Distance": "normalized_edit_distance",
            "BLEU_Score": "bleu_score",
            "ROUGE1": "rouge1",
            "ROUGE2": "rouge2",
            "ROUGEL": "rougeL",
            "LLM_Score": "LLM_Score",
            "LLM_Explanation": "LLM_Explanation",
        }
        for excel_name, internal_name in col_map.items():
            if excel_name in detailed.columns and internal_name not in detailed.columns:
                detailed[internal_name] = detailed[excel_name]
        category_analysis = pd.read_excel(
            T5_REPORT_PATH, sheet_name="Category_Analysis", engine="openpyxl"
        )
        summary_metrics = None
        try:
            summary_metrics = pd.read_excel(
                T5_REPORT_PATH, sheet_name="Summary_Metrics", engine="openpyxl"
            )
        except Exception:
            pass
        return detailed, category_analysis, summary_metrics
    except Exception as e:
        print(f"Could not load T5 report: {e}")
        return None, None, None


def run_mac_evaluation(test_df, model, evaluator):
    """Run Mac Foundation Model on test dataset and compute metrics."""
    results = []
    n_total = len(test_df)
    _write_progress("mac_model", 0, n_total)
    for idx, row in test_df.iterrows():
        input_text = row["Input_Text"]
        ground_truth = row["Ground_Truth"]
        category = row["Category"]
        all_references = row["All_References"]

        try:
            predicted = model.correct(input_text, category=category)
        except Exception as e:
            print(f"Error processing row {idx}: {e}")
            predicted = input_text

        primary_metrics = evaluator.calculate_all_metrics(predicted, ground_truth)
        multi_ref_metrics = {}
        if category == "Rephrasing" and all_references:
            multi_ref_metrics = evaluate_with_multiple_references(
                predicted, all_references, evaluator, max_refs=4
            )

        result = {
            "Index": idx,
            "Category": category,
            "Input_Text": input_text,
            "Ground_Truth": ground_truth,
            "Predicted": predicted,
            "All_References": "|".join(all_references) if all_references else "",
            **primary_metrics,
            **{k: v for k, v in multi_ref_metrics.items() if k.startswith("avg_") or k == "best_ref_score"},
        }
        results.append(result)
        completed = len(results)
        _write_progress("mac_model", completed, n_total)
        if completed % 20 == 0:
            print(f"Processed {completed}/{n_total} test cases...")

    return pd.DataFrame(results)


def build_comparison_table(mac_category_metrics, t5_detailed, t5_category_analysis, t5_summary, mac_results_df):
    """Build side-by-side comparison: T5 vs Mac Foundation Model."""
    metrics_rows = []
    if mac_category_metrics is not None and not mac_category_metrics.empty:
        for _, row in mac_category_metrics.iterrows():
            cat = row["Category"]
            metrics_rows.append({
                "Category": cat,
                "Metric": "Exact_Match_Rate",
                "T5_base_grammar_correction": _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, cat, "Exact_Match_Rate"),
                "Mac_Foundation_Model": row.get("Exact_Match_Rate"),
            })
            metrics_rows.append({
                "Category": cat,
                "Metric": "Avg_BLEU_Score",
                "T5_base_grammar_correction": _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, cat, "Avg_BLEU"),
                "Mac_Foundation_Model": row.get("Avg_BLEU_Score"),
            })
            metrics_rows.append({
                "Category": cat,
                "Metric": "Avg_ROUGEL",
                "T5_base_grammar_correction": _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, cat, "Avg_ROUGEL"),
                "Mac_Foundation_Model": row.get("Avg_ROUGEL"),
            })
            metrics_rows.append({
                "Category": cat,
                "Metric": "Avg_Normalized_Edit_Distance",
                "T5_base_grammar_correction": _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, cat, "Avg_Normalized_Edit_Distance"),
                "Mac_Foundation_Model": row.get("Avg_Normalized_Edit_Distance"),
            })
            metrics_rows.append({
                "Category": cat,
                "Metric": "High_Quality_Corrections",
                "T5_base_grammar_correction": _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, cat, "High_Quality_Corrections"),
                "Mac_Foundation_Model": _mac_high_quality(mac_results_df, cat),
            })
            if "LLM_Score" in mac_results_df.columns and mac_results_df["LLM_Score"].notna().any():
                metrics_rows.append({
                    "Category": cat,
                    "Metric": "Avg_LLM_Score",
                    "T5_base_grammar_correction": _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, cat, "Avg_LLM_Score"),
                    "Mac_Foundation_Model": _mac_llm_avg(mac_results_df, cat),
                })
    return pd.DataFrame(metrics_rows)


def _get_t5_metric(t5_category_analysis, t5_detailed, t5_summary, category, metric_key):
    if metric_key == "Avg_Normalized_Edit_Distance" and t5_detailed is not None and "normalized_edit_distance" in t5_detailed.columns:
        c = t5_detailed[t5_detailed["Category"] == category]
        return c["normalized_edit_distance"].mean() if len(c) > 0 else None
    if t5_category_analysis is None or t5_category_analysis.empty:
        return None
    row = t5_category_analysis[t5_category_analysis["Category"] == category]
    if row.empty:
        return None
    row = row.iloc[0]
    key_map = {
        "Exact_Match_Rate": "Exact_Match_Rate",
        "Avg_BLEU": "Avg_BLEU",
        "Avg_ROUGEL": "Avg_ROUGEL",
        "High_Quality_Corrections": "High_Quality_Corrections",
        "Avg_LLM_Score": "Avg_LLM_Score",
    }
    col = key_map.get(metric_key, metric_key)
    if col in row:
        return row[col]
    return None


def _mac_high_quality(mac_results_df, category):
    cat_df = mac_results_df[mac_results_df["Category"] == category]
    return len(cat_df[cat_df["normalized_edit_distance"] > 0.8]) if not cat_df.empty else 0


def _mac_llm_avg(mac_results_df, category):
    cat_df = mac_results_df[(mac_results_df["Category"] == category) & (mac_results_df["LLM_Score"].notna())]
    return cat_df["LLM_Score"].mean() if len(cat_df) > 0 else None


def build_category_comparison_summary(mac_category_metrics, t5_category_analysis, mac_results_df):
    """One row per category with T5 vs Mac columns for key metrics."""
    rows = []
    for cat in ["Typo", "Grammar", "Rephrasing", "Overall"]:
        mac_row = None
        if mac_category_metrics is not None and not mac_category_metrics.empty:
            m = mac_category_metrics[mac_category_metrics["Category"] == cat]
            mac_row = m.iloc[0] if not m.empty else None
        t5_row = None
        if t5_category_analysis is not None and not t5_category_analysis.empty:
            t = t5_category_analysis[t5_category_analysis["Category"] == cat]
            t5_row = t.iloc[0] if not t.empty else None

        def t5_val(key):
            return t5_row.get(key) if t5_row is not None else None

        def mac_val(key):
            return mac_row.get(key) if mac_row is not None else None

        mac_llm = _mac_llm_avg(mac_results_df, cat) if mac_results_df is not None and "LLM_Score" in mac_results_df.columns else None

        rows.append({
            "Category": cat,
            "T5_Exact_Match_Rate": t5_val("Exact_Match_Rate"),
            "Mac_Exact_Match_Rate": mac_val("Exact_Match_Rate"),
            "T5_Avg_BLEU": t5_val("Avg_BLEU"),
            "Mac_Avg_BLEU": mac_val("Avg_BLEU_Score"),
            "T5_Avg_ROUGEL": t5_val("Avg_ROUGEL"),
            "Mac_Avg_ROUGEL": mac_val("Avg_ROUGEL"),
            "T5_Avg_LLM_Score": t5_val("Avg_LLM_Score"),
            "Mac_Avg_LLM_Score": mac_llm,
            "T5_High_Quality_Count": t5_val("High_Quality_Corrections"),
            "Mac_High_Quality_Count": _mac_high_quality(mac_results_df, cat) if mac_results_df is not None else None,
        })
    return pd.DataFrame(rows)


def main():
    print("=" * 80)
    print("Mac Foundation Model Evaluation & Comparison with T5")
    print("=" * 80)
    print()

    if not mac_foundation_model_available() or MacFoundationModel is None:
        print("Mac Foundation Model is not available.")
        print("Requires: macOS 26+ (Tahoe), Apple Intelligence enabled, and: pip install apple-foundation-models")
        return

    # Load test dataset
    print("Step 1: Loading test dataset...")
    print("-" * 80)
    test_df = pd.read_csv(TEST_DATASET_PATH)
    test_df["All_References"] = test_df["All_References"].apply(
        lambda x: x.split("|") if isinstance(x, str) and "|" in x else ([x] if pd.notna(x) and x else [])
    )
    print(f"Loaded {len(test_df)} examples")
    print()

    # Load Mac Foundation Model
    print("Step 2: Loading Mac Foundation Model...")
    print("-" * 80)
    try:
        mac_model = MacFoundationModel()
    except Exception as e:
        print(f"Failed to load Mac Foundation Model: {e}")
        return
    print()

    # Run Mac evaluation
    print("Step 3: Evaluating Mac Foundation Model on test dataset...")
    print("Progress is written to mac_eval_progress.json (completed/total). Check it anytime.")
    print("-" * 80)
    evaluator = ModelEvaluator()
    mac_results_df = run_mac_evaluation(test_df, mac_model, evaluator)
    if hasattr(mac_model, "close"):
        mac_model.close()
    mac_category_metrics = evaluator.calculate_category_metrics(mac_results_df)
    print(mac_category_metrics.to_string(index=False))
    print()

    # LLM evaluation for Mac
    print("Step 4: Running LLM evaluation on Mac Foundation Model outputs...")
    print("-" * 80)
    n_total = len(mac_results_df)
    _write_progress("llm_eval", 0, n_total)
    try:
        llm_evaluator = LLMEvaluator()
        def _llm_progress(completed, total):
            _write_progress("llm_eval", completed, total)
        mac_results_df = llm_evaluator.evaluate_batch(
            mac_results_df,
            category_col="Category",
            input_col="Input_Text",
            truth_col="Ground_Truth",
            pred_col="Predicted",
            refs_col="All_References",
            progress_callback=_llm_progress,
        )
        print("LLM evaluation completed!")
    except Exception as e:
        print(f"LLM evaluation failed: {e}")
        mac_results_df["LLM_Score"] = None
        mac_results_df["LLM_Explanation"] = "LLM evaluation not available"
    print()

    # Load T5 results
    print("Step 5: Loading T5 evaluation results for comparison...")
    print("-" * 80)
    t5_detailed, t5_category_analysis, t5_summary = load_t5_results()
    if t5_detailed is not None:
        print(f"Loaded T5 detailed results: {len(t5_detailed)} rows")
    if t5_category_analysis is not None:
        print("Loaded T5 category analysis")
    print()

    # Build comparison tables
    print("Step 6: Building comparison tables...")
    print("-" * 80)
    comparison_table = build_comparison_table(
        mac_category_metrics, t5_detailed, t5_category_analysis, t5_summary, mac_results_df
    )
    category_comparison = build_category_comparison_summary(
        mac_category_metrics, t5_category_analysis, mac_results_df
    )
    print()

    # Write Excel
    print("Step 7: Writing Excel report...")
    print("-" * 80)

    with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
        mac_category_metrics.to_excel(writer, sheet_name="Mac_Summary_Metrics", index=False)

        detailed = mac_results_df[[
            "Index", "Category", "Input_Text", "Ground_Truth", "Predicted",
            "exact_match", "char_accuracy", "word_accuracy", "normalized_edit_distance",
            "bleu_score", "rouge1", "rouge2", "rougeL", "All_References",
        ]].copy()
        detailed.columns = [
            "Index", "Category", "Input_Text", "Ground_Truth", "Predicted",
            "Exact_Match", "Char_Accuracy", "Word_Accuracy", "Normalized_Edit_Distance",
            "BLEU_Score", "ROUGE1", "ROUGE2", "ROUGEL", "All_References",
        ]
        if "LLM_Score" in mac_results_df.columns:
            detailed["LLM_Score"] = mac_results_df["LLM_Score"]
            detailed["LLM_Explanation"] = mac_results_df["LLM_Explanation"]
        detailed.to_excel(writer, sheet_name="Mac_Detailed_Results", index=False)

        if "LLM_Score" in mac_results_df.columns and mac_results_df["LLM_Score"].notna().any():
            llm_cols = ["Index", "Category", "Input_Text", "Ground_Truth", "Predicted", "All_References", "LLM_Score", "LLM_Explanation"]
            mac_results_df[llm_cols].to_excel(writer, sheet_name="Mac_LLM_Evaluation", index=False)
            llm_summary = []
            for cat in ["Typo", "Grammar", "Rephrasing"]:
                c = mac_results_df[(mac_results_df["Category"] == cat) & (mac_results_df["LLM_Score"].notna())]
                if len(c) > 0:
                    llm_summary.append({
                        "Category": cat,
                        "Total_Evaluated": len(c),
                        "Avg_LLM_Score": c["LLM_Score"].mean(),
                        "Score_5_Count": len(c[c["LLM_Score"] == 5]),
                        "Score_4_Count": len(c[c["LLM_Score"] == 4]),
                        "Score_4+_Rate": len(c[c["LLM_Score"] >= 4]) / len(c),
                    })
            if llm_summary:
                pd.DataFrame(llm_summary).to_excel(writer, sheet_name="Mac_LLM_Summary", index=False)

        category_comparison.to_excel(writer, sheet_name="Comparison_Table", index=False)
        if not comparison_table.empty:
            comparison_table.to_excel(writer, sheet_name="Comparison_Detailed", index=False)

    format_excel_report(OUTPUT_EXCEL)

    # Clear progress file when done
    try:
        if os.path.exists(PROGRESS_FILE):
            os.remove(PROGRESS_FILE)
    except Exception:
        pass

    print()
    print("=" * 80)
    print("Evaluation Complete!")
    print("=" * 80)
    print(f"Output: {OUTPUT_EXCEL}")
    print()
    print("Comparison summary (T5 vs Mac Foundation Model):")
    print(category_comparison.to_string(index=False))


if __name__ == "__main__":
    main()
