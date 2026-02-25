"""
Optimized evaluation script for jhu-clsp/jfleg dataset
With progress saving and better error handling
"""

import pandas as pd
import os
import json
from datasets import load_dataset
from model_loader import GrammarCorrectionModel
from evaluator import ModelEvaluator
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def format_excel_report(filename):
    """Format the Excel report with styling"""
    wb = load_workbook(filename)
    
    # Define styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Format header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        # Format data rows
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        ws.freeze_panes = 'A2'
    
    wb.save(filename)
    print(f"Excel report formatted and saved to {filename}")

def load_jfleg_dataset(max_samples=None):
    """Load the JFLEG dataset from HuggingFace"""
    print("Loading JFLEG dataset from HuggingFace...")
    try:
        # Load the dataset
        print("Downloading/loading dataset (this may take a moment)...")
        dataset = load_dataset("jhu-clsp/jfleg")
        
        # Convert to pandas DataFrame
        test_data = []
        
        # Process test split (primary evaluation set)
        if 'test' in dataset:
            test_items = list(dataset['test'])
            if max_samples and len(test_items) > max_samples:
                print(f"Limiting to {max_samples} samples from test set...")
                test_items = test_items[:max_samples]
            
            for item in test_items:
                sentence = item['sentence']
                corrections = item['corrections']
                primary_correction = corrections[0] if corrections else sentence
                
                test_data.append({
                    'Input_Text': sentence,
                    'Ground_Truth': primary_correction,
                    'All_References': corrections,
                    'Num_References': len(corrections)
                })
        
        df = pd.DataFrame(test_data)
        print(f"Loaded {len(df)} examples from JFLEG dataset")
        print(f"Average number of reference corrections per example: {df['Num_References'].mean():.2f}")
        
        return df, dataset
        
    except Exception as e:
        print(f"Error loading JFLEG dataset: {e}")
        import traceback
        traceback.print_exc()
        raise

def evaluate_with_multiple_references(predicted, all_references, evaluator, max_refs=4):
    """Evaluate prediction against multiple reference corrections"""
    refs_to_eval = all_references[:max_refs] if len(all_references) > max_refs else all_references
    
    best_metrics = None
    best_score = -1
    
    for ref in refs_to_eval:
        metrics = evaluator.calculate_all_metrics(predicted, ref)
        score = metrics['normalized_edit_distance']
        
        if score > best_score:
            best_score = score
            best_metrics = metrics
    
    all_metrics = []
    for ref in refs_to_eval:
        all_metrics.append(evaluator.calculate_all_metrics(predicted, ref))
    
    avg_metrics = {}
    if all_metrics:
        for key in all_metrics[0].keys():
            avg_metrics[f'avg_{key}'] = sum(m[key] for m in all_metrics) / len(all_metrics)
    
    result_metrics = best_metrics.copy() if best_metrics else {}
    result_metrics.update(avg_metrics)
    result_metrics['best_ref_score'] = best_score
    
    return result_metrics

def evaluate_jfleg_dataset(df, model, evaluator, checkpoint_file='jfleg_progress.json'):
    """Evaluate model on JFLEG dataset with progress saving"""
    print("Evaluating model on JFLEG dataset...")
    print(f"Total examples to process: {len(df)}")
    
    # Load checkpoint if exists
    start_idx = 0
    results = []
    if os.path.exists(checkpoint_file):
        try:
            with open(checkpoint_file, 'r') as f:
                checkpoint = json.load(f)
                results = checkpoint.get('results', [])
                start_idx = checkpoint.get('last_index', 0) + 1
                print(f"Resuming from index {start_idx} (found {len(results)} completed results)")
        except:
            print("Could not load checkpoint, starting fresh")
    
    for idx in range(start_idx, len(df)):
        row = df.iloc[idx]
        input_text = row['Input_Text']
        all_references = row['All_References']
        primary_truth = row['Ground_Truth']
        
        # Get model prediction
        try:
            predicted = model.correct(input_text)
        except Exception as e:
            print(f"Error processing row {idx}: {e}")
            predicted = input_text
        
        # Calculate metrics with multiple references
        try:
            metrics = evaluate_with_multiple_references(predicted, all_references, evaluator, max_refs=4)
        except Exception as e:
            print(f"Error in multi-reference evaluation for row {idx}: {e}")
            metrics = {}
        
        # Calculate against primary reference
        try:
            primary_metrics = evaluator.calculate_all_metrics(predicted, primary_truth)
        except Exception as e:
            print(f"Error calculating primary metrics for row {idx}: {e}")
            primary_metrics = {}
        
        # Store results
        result = {
            'Index': idx,
            'Input_Text': input_text,
            'Primary_Ground_Truth': primary_truth,
            'Predicted': predicted,
            'Num_References': row['Num_References'],
            **primary_metrics,
            **{k: v for k, v in metrics.items() if k.startswith('avg_') or k == 'best_ref_score'}
        }
        results.append(result)
        
        # Progress updates and checkpoint saving
        if (idx + 1) % 10 == 0:
            print(f"Processed {idx + 1}/{len(df)} test cases ({(idx+1)/len(df)*100:.1f}%)...")
            # Save checkpoint
            try:
                with open(checkpoint_file, 'w') as f:
                    json.dump({'last_index': idx, 'results': results}, f, indent=2)
            except:
                pass
        
        # Save checkpoint every 50 examples
        if (idx + 1) % 50 == 0:
            try:
                with open(checkpoint_file, 'w') as f:
                    json.dump({'last_index': idx, 'results': results}, f, indent=2)
                print(f"Checkpoint saved at {idx + 1} examples")
            except:
                pass
    
    # Clean up checkpoint file
    if os.path.exists(checkpoint_file):
        try:
            os.remove(checkpoint_file)
        except:
            pass
    
    results_df = pd.DataFrame(results)
    return results_df

def calculate_jfleg_summary_metrics(results_df):
    """Calculate summary metrics for JFLEG evaluation"""
    summary = {
        'Metric': [],
        'Value': []
    }
    
    summary['Metric'].extend([
        'Total_Test_Cases',
        'Exact_Match_Rate (Primary)',
        'Avg_Char_Accuracy (Primary)',
        'Avg_Word_Accuracy (Primary)',
        'Avg_Normalized_Edit_Distance (Primary)',
        'Avg_BLEU_Score (Primary)',
        'Avg_ROUGE1 (Primary)',
        'Avg_ROUGE2 (Primary)',
        'Avg_ROUGEL (Primary)',
        'Exact_Matches (Primary)'
    ])
    
    summary['Value'].extend([
        len(results_df),
        results_df['exact_match'].mean(),
        results_df['char_accuracy'].mean(),
        results_df['word_accuracy'].mean(),
        results_df['normalized_edit_distance'].mean(),
        results_df['bleu_score'].mean(),
        results_df['rouge1'].mean(),
        results_df['rouge2'].mean(),
        results_df['rougeL'].mean(),
        results_df['exact_match'].sum()
    ])
    
    if 'best_ref_score' in results_df.columns:
        summary['Metric'].extend([
            'Avg_Best_Ref_Score',
            'High_Quality_Corrections (Best Ref)'
        ])
        summary['Value'].extend([
            results_df['best_ref_score'].mean(),
            (results_df['best_ref_score'] > 0.8).sum()
        ])
    
    if 'avg_normalized_edit_distance' in results_df.columns:
        summary['Metric'].extend([
            'Avg_Normalized_Edit_Distance (All Refs)',
            'Avg_BLEU_Score (All Refs)',
            'Avg_ROUGEL (All Refs)'
        ])
        summary['Value'].extend([
            results_df['avg_normalized_edit_distance'].mean(),
            results_df['avg_bleu_score'].mean(),
            results_df['avg_rougeL'].mean()
        ])
    
    return pd.DataFrame(summary)

def main():
    print("=" * 80)
    print("JFLEG Dataset Evaluation (Optimized)")
    print("=" * 80)
    print()
    
    # Step 1: Load JFLEG dataset
    print("Step 1: Loading JFLEG dataset...")
    print("-" * 80)
    jfleg_df, dataset = load_jfleg_dataset(max_samples=None)
    print()
    
    # Step 2: Load model
    print("Step 2: Loading grammar correction model...")
    print("-" * 80)
    model = GrammarCorrectionModel()
    print()
    
    # Step 3: Evaluate model
    print("Step 3: Evaluating model on JFLEG dataset...")
    print("-" * 80)
    print("Note: This may take 30-60 minutes for the full test set (748 examples)")
    print("Progress will be saved periodically...")
    print()
    evaluator = ModelEvaluator()
    results_df = evaluate_jfleg_dataset(jfleg_df, model, evaluator)
    print()
    
    # Step 4: Calculate summary metrics
    print("Step 4: Calculating summary metrics...")
    print("-" * 80)
    summary_df = calculate_jfleg_summary_metrics(results_df)
    print(summary_df.to_string(index=False))
    print()
    
    # Step 5: Create Excel report
    print("Step 5: Creating Excel evaluation report...")
    print("-" * 80)
    
    excel_filename = "jfleg_evaluation_report.xlsx"
    
    with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
        summary_df.to_excel(writer, sheet_name='Summary_Metrics', index=False)
        
        detailed_results = results_df[[
            'Index', 'Input_Text', 'Primary_Ground_Truth', 'Predicted', 'Num_References',
            'exact_match', 'char_accuracy', 'word_accuracy', 'normalized_edit_distance',
            'bleu_score', 'rouge1', 'rouge2', 'rougeL'
        ]].copy()
        
        if 'best_ref_score' in results_df.columns:
            detailed_results['best_ref_score'] = results_df['best_ref_score']
        if 'avg_normalized_edit_distance' in results_df.columns:
            detailed_results['avg_normalized_edit_distance'] = results_df['avg_normalized_edit_distance']
            detailed_results['avg_bleu_score'] = results_df['avg_bleu_score']
            detailed_results['avg_rougeL'] = results_df['avg_rougeL']
        
        column_mapping = {
            'exact_match': 'Exact_Match',
            'char_accuracy': 'Char_Accuracy',
            'word_accuracy': 'Word_Accuracy',
            'normalized_edit_distance': 'Normalized_Edit_Distance',
            'bleu_score': 'BLEU_Score',
            'rouge1': 'ROUGE1',
            'rouge2': 'ROUGE2',
            'rougeL': 'ROUGEL'
        }
        detailed_results.rename(columns=column_mapping, inplace=True)
        detailed_results.to_excel(writer, sheet_name='Detailed_Results', index=False)
        
        best_cases = results_df.nlargest(30, 'normalized_edit_distance')[
            ['Input_Text', 'Primary_Ground_Truth', 'Predicted', 'normalized_edit_distance']
        ].copy()
        best_cases.columns = ['Input_Text', 'Ground_Truth', 'Predicted', 'Score']
        best_cases.to_excel(writer, sheet_name='Best_Cases', index=False)
        
        worst_cases = results_df.nsmallest(30, 'normalized_edit_distance')[
            ['Input_Text', 'Primary_Ground_Truth', 'Predicted', 'normalized_edit_distance']
        ].copy()
        worst_cases.columns = ['Input_Text', 'Ground_Truth', 'Predicted', 'Score']
        worst_cases.to_excel(writer, sheet_name='Worst_Cases', index=False)
        
        stats_data = {
            'Statistic': [
                'Total Examples',
                'Average References per Example',
                'Min References',
                'Max References',
                'Average Input Length',
                'Average Output Length'
            ],
            'Value': [
                len(jfleg_df),
                jfleg_df['Num_References'].mean(),
                jfleg_df['Num_References'].min(),
                jfleg_df['Num_References'].max(),
                jfleg_df['Input_Text'].str.len().mean(),
                jfleg_df['Ground_Truth'].str.len().mean()
            ]
        }
        stats_df = pd.DataFrame(stats_data)
        stats_df.to_excel(writer, sheet_name='Dataset_Statistics', index=False)
    
    format_excel_report(excel_filename)
    
    print()
    print("=" * 80)
    print("JFLEG Evaluation Complete!")
    print("=" * 80)
    print(f"\nFiles generated:")
    print(f"  {excel_filename} - Comprehensive JFLEG evaluation report")
    print()
    print("Summary:")
    print(f"  Total test cases: {len(results_df)}")
    print(f"  Exact matches (primary): {results_df['exact_match'].sum()} ({results_df['exact_match'].mean()*100:.2f}%)")
    print(f"  Average BLEU score (primary): {results_df['bleu_score'].mean():.4f}")
    print(f"  Average ROUGE-L score (primary): {results_df['rougeL'].mean():.4f}")
    print(f"  Average normalized edit distance (primary): {results_df['normalized_edit_distance'].mean():.4f}")
    if 'best_ref_score' in results_df.columns:
        print(f"  Average best reference score: {results_df['best_ref_score'].mean():.4f}")

if __name__ == "__main__":
    main()
