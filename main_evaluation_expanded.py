"""
Main evaluation script with expanded dataset and LLM evaluation
Uses expanded dataset generator with 100+ examples per category and 4 references for rephrasing
Includes LLM-based evaluation on 1-5 scale
"""

import pandas as pd
from model_loader import GrammarCorrectionModel
from dataset_generator_expanded import ExpandedDatasetGenerator
from evaluator import ModelEvaluator
from llm_evaluator import LLMEvaluator
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
from dotenv import load_dotenv

# Load environment variables from .env file
load_dotenv()

def format_excel_report(filename):
    """Format the Excel report with styling"""
    wb = load_workbook(filename)
    
    # Define styles
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    category_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Format header row
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        
        # Format data rows
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                
                # Highlight category column
                if cell.column == 2:  # Category column
                    cell.fill = category_fill
                    cell.font = Font(bold=True)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        ws.freeze_panes = 'A2'
    
    wb.save(filename)
    print(f"Excel report formatted and saved to {filename}")

def evaluate_with_multiple_references(predicted, all_references, evaluator, max_refs=4):
    """Evaluate prediction against multiple reference corrections"""
    if not all_references:
        return {}
    
    refs_to_eval = all_references[:max_refs] if len(all_references) > max_refs else all_references
    
    best_metrics = None
    best_score = -1
    
    for ref in refs_to_eval:
        metrics = evaluator.calculate_all_metrics(predicted, ref)
        score = metrics['normalized_edit_distance']
        
        if score > best_score:
            best_score = score
            best_metrics = metrics
    
    all_metrics = []
    for ref in refs_to_eval:
        all_metrics.append(evaluator.calculate_all_metrics(predicted, ref))
    
    avg_metrics = {}
    if all_metrics:
        for key in all_metrics[0].keys():
            avg_metrics[f'avg_{key}'] = sum(m[key] for m in all_metrics) / len(all_metrics)
    
    result_metrics = best_metrics.copy() if best_metrics else {}
    result_metrics.update(avg_metrics)
    result_metrics['best_ref_score'] = best_score
    
    return result_metrics

def main():
    print("=" * 80)
    print("Expanded Grammar Correction Model Evaluation")
    print("=" * 80)
    print()
    
    # Step 1: Generate expanded test dataset
    print("Step 1: Generating expanded synthetic test dataset...")
    print("-" * 80)
    generator = ExpandedDatasetGenerator()
    test_df = generator.save_to_csv("test_dataset_expanded.csv")
    
    # Parse All_References from string format
    test_df['All_References'] = test_df['All_References'].apply(
        lambda x: x.split('|') if isinstance(x, str) and '|' in x else ([x] if pd.notna(x) and x else [])
    )
    print()
    
    # Step 2: Load model
    print("Step 2: Loading grammar correction model...")
    print("-" * 80)
    model = GrammarCorrectionModel()
    print()
    
    # Step 3: Evaluate model
    print("Step 3: Evaluating model on expanded test dataset...")
    print("-" * 80)
    evaluator = ModelEvaluator()
    results = []
    
    for idx, row in test_df.iterrows():
        input_text = row['Input_Text']
        ground_truth = row['Ground_Truth']
        category = row['Category']
        all_references = row['All_References']
        
        # Get model prediction
        try:
            predicted = model.correct(input_text)
        except Exception as e:
            print(f"Error processing row {idx}: {e}")
            predicted = input_text
        
        # Calculate metrics
        primary_metrics = evaluator.calculate_all_metrics(predicted, ground_truth)
        
        # For rephrasing, also calculate against all references
        multi_ref_metrics = {}
        if category == 'Rephrasing' and all_references:
            multi_ref_metrics = evaluate_with_multiple_references(predicted, all_references, evaluator, max_refs=4)
        
        # Store results
        result = {
            'Index': idx,
            'Category': category,
            'Input_Text': input_text,
            'Ground_Truth': ground_truth,
            'Predicted': predicted,
            'All_References': '|'.join(all_references) if all_references else '',
            **primary_metrics,
            **{k: v for k, v in multi_ref_metrics.items() if k.startswith('avg_') or k == 'best_ref_score'}
        }
        results.append(result)
        
        if (idx + 1) % 20 == 0:
            print(f"Processed {idx + 1}/{len(test_df)} test cases...")
    
    results_df = pd.DataFrame(results)
    print()
    
    # Step 4: Calculate aggregated metrics
    print("Step 4: Calculating aggregated metrics...")
    print("-" * 80)
    category_metrics_df = evaluator.calculate_category_metrics(results_df)
    print(category_metrics_df.to_string(index=False))
    print()
    
    # Step 5: LLM Evaluation
    print("Step 5: Running LLM-based evaluation (1-5 scale)...")
    print("-" * 80)
    print("Note: This requires API credentials in .env file")
    print("  For Azure OpenAI: AZURE_OPENAI_API_KEY, AZURE_OPENAI_ENDPOINT, AZURE_OPENAI_DEPLOYMENT_NAME")
    print("  For OpenAI: OPENAI_API_KEY")
    print("This may take several minutes and incur API costs...")
    
    try:
        # LLMEvaluator will automatically detect Azure or OpenAI from .env file
        llm_evaluator = LLMEvaluator()
        results_df = llm_evaluator.evaluate_batch(
            results_df,
            category_col='Category',
            input_col='Input_Text',
            truth_col='Ground_Truth',
            pred_col='Predicted',
            refs_col='All_References'
        )
        print("LLM evaluation completed!")
    except Exception as e:
        print(f"LLM evaluation failed: {e}")
        print("Continuing without LLM scores...")
        results_df['LLM_Score'] = None
        results_df['LLM_Explanation'] = 'LLM evaluation not available'
    print()
    
    # Step 6: Create Excel report
    print("Step 6: Creating Excel evaluation report...")
    print("-" * 80)
    
    excel_filename = "model_evaluation_expanded_report.xlsx"
    
    with pd.ExcelWriter(excel_filename, engine='openpyxl') as writer:
        # Sheet 1: Summary Metrics by Category
        category_metrics_df.to_excel(writer, sheet_name='Summary_Metrics', index=False)
        
        # Sheet 2: Detailed Results
        detailed_results = results_df[[
            'Index', 'Category', 'Input_Text', 'Ground_Truth', 'Predicted',
            'exact_match', 'char_accuracy', 'word_accuracy', 'normalized_edit_distance',
            'bleu_score', 'rouge1', 'rouge2', 'rougeL'
        ]].copy()
        
        # Add All_References column if available (important for rephrasing with 4 references)
        if 'All_References' in results_df.columns:
            detailed_results['All_References'] = results_df['All_References']
        
        # Add LLM scores if available
        if 'LLM_Score' in results_df.columns:
            detailed_results['LLM_Score'] = results_df['LLM_Score']
            detailed_results['LLM_Explanation'] = results_df['LLM_Explanation']
        
        # Rename columns
        column_mapping = {
            'exact_match': 'Exact_Match',
            'char_accuracy': 'Char_Accuracy',
            'word_accuracy': 'Word_Accuracy',
            'normalized_edit_distance': 'Normalized_Edit_Distance',
            'bleu_score': 'BLEU_Score',
            'rouge1': 'ROUGE1',
            'rouge2': 'ROUGE2',
            'rougeL': 'ROUGEL'
        }
        detailed_results.rename(columns=column_mapping, inplace=True)
        
        detailed_results.to_excel(writer, sheet_name='Detailed_Results', index=False)
        
        # Sheet 3: Best and Worst Cases
        best_cases = results_df.nlargest(20, 'normalized_edit_distance')[
            ['Category', 'Input_Text', 'Ground_Truth', 'Predicted', 'normalized_edit_distance']
        ].copy()
        best_cases.columns = ['Category', 'Input_Text', 'Ground_Truth', 'Predicted', 'Score']
        best_cases.to_excel(writer, sheet_name='Best_Cases', index=False)
        
        worst_cases = results_df.nsmallest(20, 'normalized_edit_distance')[
            ['Category', 'Input_Text', 'Ground_Truth', 'Predicted', 'normalized_edit_distance']
        ].copy()
        worst_cases.columns = ['Category', 'Input_Text', 'Ground_Truth', 'Predicted', 'Score']
        worst_cases.to_excel(writer, sheet_name='Worst_Cases', index=False)
        
        # Sheet 4: Category-wise Analysis
        category_analysis = []
        for category in ['Typo', 'Grammar', 'Rephrasing']:
            cat_df = results_df[results_df['Category'] == category]
            if len(cat_df) > 0:
                exact_matches = cat_df[cat_df['exact_match'] == True]
                analysis = {
                    'Category': category,
                    'Total_Cases': len(cat_df),
                    'Exact_Matches': len(exact_matches),
                    'Exact_Match_Rate': len(exact_matches) / len(cat_df),
                    'Avg_BLEU': cat_df['bleu_score'].mean(),
                    'Avg_ROUGE1': cat_df['rouge1'].mean(),
                    'Avg_ROUGEL': cat_df['rougeL'].mean(),
                    'High_Quality_Corrections': len(cat_df[cat_df['normalized_edit_distance'] > 0.8])
                }
                # Add LLM metrics if available
                if 'LLM_Score' in cat_df.columns and cat_df['LLM_Score'].notna().any():
                    analysis['Avg_LLM_Score'] = cat_df['LLM_Score'].mean()
                    analysis['LLM_Score_5'] = len(cat_df[cat_df['LLM_Score'] == 5])
                    analysis['LLM_Score_4+'] = len(cat_df[cat_df['LLM_Score'] >= 4])
                
                category_analysis.append(analysis)
        
        category_analysis_df = pd.DataFrame(category_analysis)
        category_analysis_df.to_excel(writer, sheet_name='Category_Analysis', index=False)
        
        # Sheet 5: LLM Evaluation Results
        if 'LLM_Score' in results_df.columns and results_df['LLM_Score'].notna().any():
            llm_cols = ['Index', 'Category', 'Input_Text', 'Ground_Truth', 'Predicted', 'LLM_Score', 'LLM_Explanation']
            # Add All_References if available (important for rephrasing)
            if 'All_References' in results_df.columns:
                llm_cols.insert(5, 'All_References')  # Insert after Predicted
            
            llm_results = results_df[llm_cols].copy()
            llm_results.to_excel(writer, sheet_name='LLM_Evaluation', index=False)
            
            # LLM Summary by Category
            llm_summary = []
            for category in ['Typo', 'Grammar', 'Rephrasing']:
                cat_df = results_df[(results_df['Category'] == category) & (results_df['LLM_Score'].notna())]
                if len(cat_df) > 0:
                    llm_summary.append({
                        'Category': category,
                        'Total_Evaluated': len(cat_df),
                        'Avg_LLM_Score': cat_df['LLM_Score'].mean(),
                        'Score_5_Count': len(cat_df[cat_df['LLM_Score'] == 5]),
                        'Score_4_Count': len(cat_df[cat_df['LLM_Score'] == 4]),
                        'Score_3_Count': len(cat_df[cat_df['LLM_Score'] == 3]),
                        'Score_2_Count': len(cat_df[cat_df['LLM_Score'] == 2]),
                        'Score_1_Count': len(cat_df[cat_df['LLM_Score'] == 1]),
                        'Score_4+_Rate': len(cat_df[cat_df['LLM_Score'] >= 4]) / len(cat_df)
                    })
            
            if llm_summary:
                llm_summary_df = pd.DataFrame(llm_summary)
                llm_summary_df.to_excel(writer, sheet_name='LLM_Summary', index=False)
    
    # Format the Excel file
    format_excel_report(excel_filename)
    
    print()
    print("=" * 80)
    print("Evaluation Complete!")
    print("=" * 80)
    print(f"\nFiles generated:")
    print(f"  1. test_dataset_expanded.csv - Expanded test dataset with ground truth")
    print(f"  2. {excel_filename} - Comprehensive evaluation report")
    print()
    print("Summary:")
    print(f"  Total test cases: {len(results_df)}")
    print(f"  Exact matches: {results_df['exact_match'].sum()} ({results_df['exact_match'].mean()*100:.2f}%)")
    print(f"  Average BLEU score: {results_df['bleu_score'].mean():.4f}")
    print(f"  Average ROUGE-L score: {results_df['rougeL'].mean():.4f}")
    print(f"  Average normalized edit distance: {results_df['normalized_edit_distance'].mean():.4f}")
    if 'LLM_Score' in results_df.columns and results_df['LLM_Score'].notna().any():
        print(f"  Average LLM score: {results_df['LLM_Score'].mean():.2f}/5.0")

if __name__ == "__main__":
    main()
